"""Lógica de negocio del STATUS por SO o por shipment.

Lee los Excel crudos (ventas o shipment, y stock), cruza contra producción,
detecta producto retenido y arma el resultado que muestra la web.
No guarda nada en disco.
"""
import io
import json
import re
from pathlib import Path

import pandas as pd

_RAIZ = Path(__file__).parent
CONFIG_DIR = _RAIZ if (_RAIZ / "columnas.json").exists() else _RAIZ / "config"

# Columnas del Excel de shipment (se pueden redefinir en columnas.json, clave "shipment").
# Cada dato acepta varios nombres posibles de encabezado.
SHIPMENT_POR_DEFECTO = {
    "so": ["Nro. Confir.", "Order Number"],
    "pedidos": ["P. Twins"],
    "codigo": ["Cod.", "Code"],
    "nivel3": ["Cod Nvl 3.", "Cod Nvl 3"],
    "descripcion": ["Desc.", "Description"],
    "kg": ["Cant.", "Quantity"],
    "consignee": ["Desc. Cli.", "Consignee"],
}


class StatusError(Exception):
    """Error que se muestra tal cual a quien usa la app."""


# ---------------------------------------------------------------- configuración
def cargar_config():
    def _leer(nombre):
        with open(CONFIG_DIR / nombre, encoding="utf-8") as f:
            return json.load(f)

    cols, reglas, equivalencias = _leer("columnas.json"), _leer("reglas.json"), _leer("equivalencias.json")
    cols.setdefault("shipment", SHIPMENT_POR_DEFECTO)
    cols["stock"].setdefault("pedido", "PEDIDO")
    return cols, reglas, equivalencias


# ---------------------------------------------------------------- normalización
def norm_so(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor != valor:  # NaN
        return ""
    return re.sub(r"\.0+$", "", str(valor).strip())


def norm_codigo(valor):
    return str(valor).strip().upper()


def lista_pedidos(texto):
    return [p for p in (norm_so(x) for x in re.split(r"[,;/\s]+", str(texto or ""))) if p]


def limpiar_descripcion(desc, prefijos):
    texto = re.sub(r"\s+", " ", str(desc)).strip()
    texto = re.sub(r"\s*\(\d+\s*DAYS\)\s*$", "", texto, flags=re.I)
    for p in prefijos:
        p = re.sub(r"\s+", " ", p).strip()
        if texto.upper().startswith(p.upper()):
            resto = texto[len(p):].strip()
            if resto:
                return resto
    return texto or str(desc)


# ---------------------------------------------------------------- formato es-AR
def fmt_num(x):
    if abs(x - round(x)) < 0.005:
        s = f"{round(x):,}"
    else:
        s = f"{x:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_pct(p):
    if p is None:
        return "s/SC"
    v = p * 100
    if abs(v) < 0.005:
        return "0%"
    s = f"{v:+.2f}" if abs(v) < 0.05 else f"{v:+.1f}"
    s = s.replace(".", ",")
    if s.endswith(",0"):
        s = s[:-2]
    return s + "%"


def texto_estado(estado, pct):
    p = fmt_pct(pct)
    if estado == "ok":
        return f"🟢 {p}"
    if estado == "pend":
        return f"🔴🚨 {p}"
    return f"🟠🚨 Excede about {p}"


# ---------------------------------------------------------------- lectura de Excel
def _abrir_hojas(data, etiqueta):
    """Devuelve una función por hoja que entrega sus filas. Usa python-calamine
    (muy rápido) y, si no está instalado, openpyxl."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        CalamineWorkbook = None

    try:
        if CalamineWorkbook is not None:
            wb = CalamineWorkbook.from_filelike(io.BytesIO(data))
            return [(lambda n=n: wb.get_sheet_by_name(n).iter_rows()) for n in wb.sheet_names]
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [(lambda ws=ws: ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    except Exception as exc:
        raise StatusError(
            f"No se pudo abrir el Excel de {etiqueta}. Verifique que sea un archivo .xlsx válido."
        ) from exc


def _alternativas(valor):
    return [valor] if isinstance(valor, str) else list(valor)


def _leer_tabla(data, mapa, requeridas, opcionales, etiqueta, filtro_col=None, filtro_vals=None):
    """Recorre el Excel una sola vez, detecta la fila de encabezados (puede haber
    títulos arriba) y conserva solo las columnas configuradas. Si se indica
    filtro_col, se quedan solo las filas cuyo valor esté en filtro_vals."""
    aceptadas = {
        k: [a.strip().lower() for a in _alternativas(mapa[k])] for k in requeridas + opcionales if k in mapa
    }
    esperadas = [_alternativas(mapa[k])[0] for k in requeridas]

    filas = []
    posiciones = None
    filtro_aplicado = False
    vista = []
    for abrir in _abrir_hojas(data, etiqueta):
        it = iter(abrir())
        for n_fila, fila in enumerate(it):
            if n_fila >= 40:
                break
            celdas = {}
            for i, v in enumerate(fila):
                if v not in (None, ""):
                    celdas.setdefault(str(v).strip().lower(), i)
            textos = [str(v).strip() for v in fila if v not in (None, "")]
            if len(textos) > len(vista):
                vista = textos
            pos = {}
            for k, alternativas in aceptadas.items():
                for a in alternativas:
                    if a in celdas:
                        pos[k] = celdas[a]
                        break
            if all(k in pos for k in requeridas):
                posiciones = pos
                break
        if posiciones is None:
            continue
        idx = posiciones.get(filtro_col) if filtro_col else None
        filtro_aplicado = idx is not None and filtro_vals is not None
        for fila in it:
            if filtro_aplicado and not (idx < len(fila) and norm_so(fila[idx]) in filtro_vals):
                continue
            filas.append({k: (fila[i] if i < len(fila) else None) for k, i in posiciones.items()})
        break

    if posiciones is None:
        raise StatusError(
            f"No se encontraron las columnas esperadas en el Excel de {etiqueta}: "
            f"{', '.join(esperadas)}. Columnas detectadas: {', '.join(vista[:14]) or 'ninguna'}. "
            f"Si cambiaron de nombre, se ajustan en columnas.json."
        )
    df = pd.DataFrame(filas, columns=list(posiciones))
    df.attrs["filtro_aplicado"] = filtro_aplicado
    return df


# ---------------------------------------------------------------- líneas de venta
def _lineas_de_ventas(ventas, so):
    v = ventas.copy()
    if v.empty:
        raise StatusError(f"No se encontró el SO {so} en el Excel de ventas.")
    v["codigo"] = v["codigo"].map(norm_codigo)
    v["kg"] = pd.to_numeric(v["kg"], errors="coerce").fillna(0.0)
    agrupado = v.groupby("codigo", sort=False).agg(desc=("descripcion", "first"), sc=("kg", "sum")).reset_index()
    lineas = [{"codigo": r["codigo"], "desc": r["desc"], "sc": float(r["sc"])} for _, r in agrupado.iterrows()]
    consignee = ""
    if "consignee" in v.columns:
        no_nulos = v["consignee"].dropna()
        consignee = str(no_nulos.iloc[0]).strip() if not no_nulos.empty else ""
    return lineas, {}, consignee, []


def _lineas_de_shipment(sh):
    """Cada línea del shipment puede venir repetida una vez por código de nivel 3
    (misma cantidad). Si hay varios códigos de nivel 3 con la misma cantidad, es una
    sola línea; el nivel 3 sirve además como equivalencia automática con producción."""
    s = sh.copy()
    s["codigo"] = s["codigo"].map(norm_codigo)
    s["kg"] = pd.to_numeric(s["kg"], errors="coerce").fillna(0.0)

    lineas, equiv = [], {}
    for codigo, g in s.groupby("codigo", sort=False):
        kgs = g["kg"].tolist()
        nivel3 = []
        if "nivel3" in g.columns:
            for x in g["nivel3"].tolist():
                c = norm_codigo(x) if norm_so(x) else ""
                if c and c not in ("0", "NAN", "NONE") and c not in nivel3:
                    nivel3.append(c)
        if len(nivel3) > 1 and len({round(k, 6) for k in kgs}) == 1:
            sc = float(kgs[0])
        else:
            sc = float(sum(kgs))
        if nivel3:
            equiv[codigo] = nivel3
        desc = g["descripcion"].iloc[0] if "descripcion" in g.columns else codigo
        lineas.append({"codigo": codigo, "desc": desc, "sc": sc})

    consignee = ""
    if "consignee" in s.columns:
        no_nulos = s["consignee"].dropna()
        consignee = str(no_nulos.iloc[0]).strip() if not no_nulos.empty else ""
    pedidos = []
    if "pedidos" in s.columns:
        for celda in s["pedidos"].dropna().tolist():
            for p in lista_pedidos(celda):
                if p not in pedidos:
                    pedidos.append(p)
    return lineas, equiv, consignee, pedidos


# ---------------------------------------------------------------- cálculo
def clasificar(sc, producido, tolerancia):
    diff = producido - sc
    if sc <= 0:
        return diff, None, ("exc" if producido > 0 else "ok")
    pct = diff / sc
    if abs(pct) <= tolerancia + 1e-9:
        return diff, pct, "ok"
    return diff, pct, ("pend" if diff < 0 else "exc")


def generar_status(ventas_bytes, stock_bytes, so, shipment_bytes=None, pedidos_txt=""):
    cols, reglas, equivalencias = cargar_config()
    tolerancia = float(reglas.get("tolerancia", 0.10))
    aptos = {e.strip().upper() for e in reglas.get("estados_aptos", ["APTO"])}
    prefijos = reglas.get("prefijos_descripcion", [])
    so = norm_so(so)
    if not so:
        raise StatusError("Ingrese el número de SO o de shipment.")

    es_shipment = shipment_bytes is not None
    pedidos = []

    if es_shipment:
        sh = _leer_tabla(
            shipment_bytes, cols["shipment"], ["codigo", "kg"],
            ["so", "pedidos", "descripcion", "nivel3", "consignee"], "shipment", "so", {so},
        )
        if sh.empty:
            if sh.attrs.get("filtro_aplicado"):
                raise StatusError(f"No se encontró el shipment {so} en el Excel de shipment.")
            raise StatusError("El Excel de shipment no tiene líneas.")
        lineas_venta, equiv_auto, consignee, pedidos_auto = _lineas_de_shipment(sh)
        pedidos = lista_pedidos(pedidos_txt) or pedidos_auto
        if not pedidos:
            raise StatusError(
                "Ingrese los pedidos del shipment (por ejemplo 218136, 218137) para buscarlos en el stock."
            )
        stock = _leer_tabla(
            stock_bytes, cols["stock"], ["pedido", "codigo", "kg"], ["descripcion", "retenido"],
            "stock", "pedido", set(pedidos),
        )
    else:
        ventas = _leer_tabla(
            ventas_bytes, cols["ventas"], ["so", "codigo", "descripcion", "kg"], ["consignee"],
            "ventas", "so", {so},
        )
        lineas_venta, equiv_auto, consignee, _ = _lineas_de_ventas(ventas, so)
        stock = _leer_tabla(
            stock_bytes, cols["stock"], ["so", "codigo", "kg"], ["descripcion", "retenido"],
            "stock", "so", {so},
        )

    # --- stock del SO / shipment
    s = stock.copy()
    s["codigo"] = s["codigo"].map(norm_codigo)
    s["kg"] = pd.to_numeric(s["kg"], errors="coerce").fillna(0.0)
    if "retenido" in s.columns:
        estado = s["retenido"].fillna("").astype(str).str.strip().str.upper()
        s["ret"] = (estado != "") & (~estado.isin(aptos))
    else:
        s["ret"] = False

    prod = s.groupby("codigo")["kg"].sum().to_dict()
    ret = s[s["ret"]].groupby("codigo")["kg"].sum().to_dict()
    desc_stock = {}
    if "descripcion" in s.columns:
        desc_stock = s.groupby("codigo")["descripcion"].first().to_dict()

    def _norm_equiv(d):
        return {norm_codigo(k): [norm_codigo(c) for c in lst] for k, lst in d.items()}

    equiv = _norm_equiv(equivalencias.get("*", {}))
    equiv.update(equiv_auto)
    madre = so.split("-")[0]
    if madre != so:
        equiv.update(_norm_equiv(equivalencias.get(madre, {})))
    equiv.update(_norm_equiv(equivalencias.get(so, {})))

    # --- líneas
    usados = set()
    lineas = []
    for l in lineas_venta:
        codigo = l["codigo"]
        codigos_prod = equiv.get(codigo, [codigo])
        usados.update(codigos_prod)
        producido = float(sum(prod.get(c, 0.0) for c in codigos_prod))
        retenido = float(sum(ret.get(c, 0.0) for c in codigos_prod))
        sc = float(l["sc"])
        diff, pct, est = clasificar(sc, producido, tolerancia)
        lineas.append(
            {
                "codigo": codigo,
                "producto": limpiar_descripcion(l["desc"], prefijos),
                "sc": sc,
                "producido": producido,
                "retenido": retenido,
                "diff": float(diff),
                "pct": None if pct is None else float(pct),
                "estado": est,
                "sc_txt": fmt_num(sc),
                "prod_txt": fmt_num(producido),
                "ret_txt": fmt_num(retenido),
                "diff_txt": fmt_num(diff),
                "status": texto_estado(est, pct),
            }
        )

    # --- producción sin línea de venta
    sin_venta = []
    for codigo, kg in prod.items():
        if codigo not in usados:
            sin_venta.append(
                {
                    "codigo": codigo,
                    "descripcion": str(desc_stock.get(codigo, "")),
                    "kg_txt": fmt_num(float(kg)),
                    "ret_txt": fmt_num(float(ret.get(codigo, 0.0))),
                }
            )

    # --- totales y alertas
    tot_sc = sum(l["sc"] for l in lineas)
    tot_prod = sum(l["producido"] for l in lineas)
    tot_ret = sum(l["retenido"] for l in lineas)
    ret_total = float(s.loc[s["ret"], "kg"].sum())
    ret_codigos = int(s.loc[s["ret"], "codigo"].nunique())
    hay_retenido = ret_total > 0.005

    pendientes = [l for l in lineas if l["estado"] == "pend"]
    excesos = [l for l in lineas if l["estado"] == "exc"]
    ok = (not hay_retenido) and not pendientes and not excesos

    resumen_pend = ""
    if pendientes:
        partes = [f"{l['producto']} ({fmt_num(-l['diff'])} kg)" for l in pendientes]
        resumen_pend = "Pendiente real: " + ", ".join(partes) + "."

    avisos = []
    if s.empty:
        if es_shipment:
            avisos.append(
                f"No hay stock cargado con PEDIDO = {', '.join(pedidos)}. Revise que los pedidos del shipment estén bien escritos."
            )
        else:
            avisos.append(
                f"No hay stock cargado con PED. ORIG = {so}. Si la producción está en otros códigos, declare la equivalencia en equivalencias.json."
            )
    if sin_venta:
        avisos.append("Producción en códigos sin línea de venta: " + ", ".join(x["codigo"] for x in sin_venta) + ".")

    return {
        "so": so,
        "consignee": consignee,
        "es_shipment": es_shipment,
        "pedidos": pedidos,
        "lineas": lineas,
        "con_retenido": tot_ret > 0.005,
        "tot": {
            "sc_txt": fmt_num(tot_sc),
            "prod_txt": fmt_num(tot_prod),
            "ret_txt": fmt_num(tot_ret),
            "diff_txt": fmt_num(tot_prod - tot_sc),
            "sc": tot_sc,
            "producido": tot_prod,
            "retenido": tot_ret,
        },
        "hay_retenido": hay_retenido,
        "retenido_txt": (
            f"{fmt_num(ret_total)} kg retenidos en {ret_codigos} código" + ("s" if ret_codigos != 1 else "")
        )
        if hay_retenido
        else "",
        "ok": ok,
        "resumen_pend": resumen_pend,
        "sin_venta": sin_venta,
        "avisos": avisos,
    }
